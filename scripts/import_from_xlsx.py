"""Import SPEC06 regression data from an Excel workbook."""

import argparse
import calendar
from dataclasses import dataclass
import logging
from pathlib import Path
import re
import time
from typing import Any, Literal

from openpyxl import load_workbook

from modules.github import GitHub
from modules.json import MetadataJson, RunListJson, ReportRegressionJson

OWNER = "OpenXiangShan"
REPO = "XiangShan"
WORKSHEET = "kmh-v3 spec06"
DATA_PATH = Path(__file__).parent.parent / "data"

Subset = Literal["spec06/gcc", "spec06/xscc"]

# The rows containing the benchmark names follow each of these markers.  The
# parser still discovers the end of a block from the next marker, so changes
# to the number of benchmark rows do not require changing these values.
BLOCK_NAMES = ("GCC16", "XSCC", "GCC15", "GCC12")
GCC_PRIORITY = ("GCC16", "GCC15", "GCC12")
BENCHMARK_RE = re.compile(r"^\d+\.[A-Za-z0-9_]+$")


@dataclass(frozen=True)
class XlsxRecord:
    """Performance values belonging to one workbook column and compiler."""

    column: int
    commit: str
    source_date: Any
    block: str
    scores: dict[str, float]


@dataclass(frozen=True)
class CommitInfo:
    """Commit metadata fetched from GitHub."""

    sha: str
    title: str
    date: int


def _cell_value(value: Any) -> Any:
    """Normalize values read from openpyxl cells."""
    if isinstance(value, str):
        value = value.strip()
    return value


def _is_score(value: Any) -> bool:
    """Return whether a workbook cell contains a usable numeric score."""
    if value is None or isinstance(value, bool):
        return False
    if isinstance(value, str):
        if not value or value.startswith("#"):
            return False
        try:
            float(value)
        except ValueError:
            return False
        return True
    return isinstance(value, (int, float))


def _find_blocks(worksheet: Any) -> dict[str, tuple[int, int]]:
    """Find compiler block ranges from their labels in column A."""
    starts: dict[str, int] = {}
    for row in range(1, worksheet.max_row + 1):
        value = _cell_value(worksheet.cell(row=row, column=1).value)
        if isinstance(value, str) and value.upper() in BLOCK_NAMES:
            starts[value.upper()] = row

    missing = [name for name in BLOCK_NAMES if name not in starts]
    if missing:
        raise ValueError(
            f"Missing compiler block(s) in {WORKSHEET}: {', '.join(missing)}"
        )

    ordered = sorted(starts.items(), key=lambda item: item[1])
    return {
        name: (start + 1, next_start - 1 if next_start else worksheet.max_row)
        for (name, start), next_start in zip(
            ordered,
            [item[1] for item in ordered[1:]] + [None],
        )
    }


def _read_block(
    worksheet: Any,
    block: str,
    row_range: tuple[int, int],
    column: int,
) -> dict[str, float]:
    """Read benchmark scores for one block and workbook column."""
    scores = {}
    for row in range(row_range[0], row_range[1] + 1):
        benchmark = _cell_value(worksheet.cell(row=row, column=1).value)
        if not isinstance(benchmark, str) or not BENCHMARK_RE.fullmatch(benchmark):
            continue
        value = _cell_value(worksheet.cell(row=row, column=column).value)
        if _is_score(value):
            scores[benchmark] = float(value)

    if not scores:
        logging.debug("No benchmark scores in %s column %s", block, column)
    return scores


def read_records(path: Path) -> tuple[list[XlsxRecord], list[XlsxRecord]]:
    """Read the target worksheet and return GCC and XSCC records."""
    if not path.is_file():
        raise ValueError(f"Invalid Excel workbook: {path}")

    # This workbook has a worksheet dimension of A1 because of its merged
    # formatting rows.  Read-only mode trusts that dimension and would hide
    # the actual rows, so use the regular worksheet reader here.
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        if WORKSHEET not in workbook.sheetnames:
            raise ValueError(f"Worksheet not found: {WORKSHEET}")
        worksheet = workbook[WORKSHEET]
        blocks = _find_blocks(worksheet)

        candidates: dict[str, dict[int, XlsxRecord]] = {
            name: {} for name in BLOCK_NAMES
        }
        header_columns = [
            column
            for column in range(2, worksheet.max_column + 1)
            if worksheet.cell(row=2, column=column).value is not None
            or worksheet.cell(row=3, column=column).value is not None
        ]
        for column in header_columns:
            source_date = _cell_value(worksheet.cell(row=2, column=column).value)
            commit = _cell_value(worksheet.cell(row=3, column=column).value)
            if source_date is None or commit is None or commit == "":
                continue
            if not isinstance(commit, str):
                commit = str(commit)

            for block, row_range in blocks.items():
                scores = _read_block(worksheet, block, row_range, column)
                if scores:
                    candidates[block][column] = XlsxRecord(
                        column=column,
                        commit=commit,
                        source_date=source_date,
                        block=block,
                        scores=scores,
                    )

        gcc_records = []
        for column in sorted(
            {column for block in GCC_PRIORITY for column in candidates[block]}
        ):
            selected = next(
                (
                    candidates[block][column]
                    for block in GCC_PRIORITY
                    if column in candidates[block]
                ),
                None,
            )
            if selected is not None:
                gcc_records.append(selected)

        xscc_records = [
            candidates["XSCC"][column] for column in sorted(candidates["XSCC"])
        ]
        return gcc_records, xscc_records
    finally:
        workbook.close()


def _commit_timestamp(commit: dict) -> int:
    """Convert GitHub's commit timestamp to the metadata timestamp format."""
    date = commit["commit"]["committer"]["date"]
    return int(calendar.timegm(time.strptime(date, "%Y-%m-%dT%H:%M:%SZ")))


def _next_run_id(metadata: MetadataJson) -> str:
    """Return the first unused imported run ID."""
    run_id = 1
    while f"imported-{run_id}" in metadata.data:
        run_id += 1
    return f"imported-{run_id}"


def _hash_for_record(metadata: MetadataJson, run_list: RunListJson, commit: str) -> str:
    """Choose a report hash, preserving duplicate records for one commit."""
    used = {metadata.data[run_id].hash for run_id in run_list.runs}
    if commit not in used:
        return commit
    suffix = 1
    while f"{commit}-{suffix}" in used:
        suffix += 1
    return f"{commit}-{suffix}"


def get_commit_info(
    gh: GitHub,
    commit_ref: str,
    cache: dict[str, CommitInfo | None],
) -> CommitInfo | None:
    """Resolve a workbook commit reference to GitHub commit metadata."""
    if commit_ref in cache:
        return cache[commit_ref]

    try:
        commit = gh.commits.get_commit(OWNER, REPO, commit_ref)
        info = CommitInfo(
            sha=commit["sha"],
            title=commit["commit"]["message"].splitlines()[0],
            date=_commit_timestamp(commit),
        )
    except (KeyError, TypeError, ValueError) as error:
        logging.warning("Invalid GitHub response for commit %s: %s", commit_ref, error)
        cache[commit_ref] = None
        return None

    cache[commit_ref] = info
    return info


def import_subset(
    gh: GitHub,
    records: list[XlsxRecord],
    subset: Subset,
    data_path: Path,
    branch: str,
    dry_run: bool,
    commit_cache: dict[str, CommitInfo | None],
    metadata: MetadataJson,
) -> None:
    """Import records into one weekly regression subset."""
    branch_path = data_path / "weekly" / "xs" / branch
    subset_path = branch_path / subset
    run_list = RunListJson.from_json(subset_path / "list.json")

    for record in records:
        info = get_commit_info(gh, record.commit, commit_cache)
        if info is None:
            continue
        if run_list.exists(metadata, info.sha):
            logging.info("Skipping existing commit %s", info.sha)
            continue

        run_id = _next_run_id(metadata)
        hash_name = _hash_for_record(metadata, run_list, info.sha)
        logging.info(
            "Importing %s column %s (%s) into %s as run %s",
            info.sha,
            record.column,
            record.block,
            subset,
            run_id,
        )
        metadata.add(run_id, hash_name, info.title, info.date)
        run_list.add(run_id)
        if not dry_run:
            report = ReportRegressionJson()
            for benchmark, score in record.scores.items():
                report.append(benchmark, score)
            report.to_json(subset_path / f"{hash_name}.json")

    if not dry_run:
        metadata.to_json(branch_path / "metadata.json")
        run_list.to_json(subset_path / "list.json")


def main() -> None:
    """Main function."""
    parser = argparse.ArgumentParser(
        description="Import SPEC06 regression data from an Excel workbook"
    )
    parser.add_argument("--token", help="GitHub personal access token", required=True)
    parser.add_argument(
        "--xlsx",
        type=Path,
        default=Path(__file__).with_name("data.xlsx"),
        help="Path to the Excel workbook",
    )
    parser.add_argument(
        "--branch",
        default="kunminghu-v3",
        help="Dashboard branch to update (default: kunminghu-v3)",
    )
    parser.add_argument(
        "--logging-level",
        default="INFO",
        help="Logging level",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Parse without writing JSON files",
    )
    args = parser.parse_args()
    logging.basicConfig(level=getattr(logging, args.logging_level.upper()))

    gcc_records, xscc_records = read_records(args.xlsx)
    logging.info(
        "Read %d GCC and %d XSCC records from %s",
        len(gcc_records),
        len(xscc_records),
        args.xlsx,
    )

    gh = GitHub(args.token)
    commit_cache: dict[str, CommitInfo | None] = {}
    metadata = MetadataJson.from_json(
        DATA_PATH / "weekly" / "xs" / args.branch / "metadata.json"
    )
    import_subset(
        gh,
        gcc_records,
        "spec06/gcc",
        DATA_PATH,
        args.branch,
        args.dry_run,
        commit_cache,
        metadata,
    )
    import_subset(
        gh,
        xscc_records,
        "spec06/xscc",
        DATA_PATH,
        args.branch,
        args.dry_run,
        commit_cache,
        metadata,
    )


if __name__ == "__main__":
    main()
